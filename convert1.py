import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
import os

xml_path = r"slides_data.xml"

try:
    tree = ET.parse(xml_path)
    root = tree.getroot()

    prs = Presentation()

    # Iterate over each slide in the XML
    for slide_elem in root.findall('slide'):
        slide_id = slide_elem.get('id')
        title = slide_elem.find('title').text
        subtitle = slide_elem.find('subtitle').text
        body = slide_elem.find('body')
        diagram_code = slide_elem.find('diagram').text.strip() if slide_elem.find('diagram') is not None else None

        # Add a title and content slide
        slide_layout = prs.slide_layouts[1]  # 1 is Title and Content layout
        slide = prs.slides.add_slide(slide_layout)

        # Set title
        title_shape = slide.shapes.title
        title_shape.text = title

        # Set subtitle in content placeholder
        content_placeholder = slide.placeholders[1]
        tf = content_placeholder.text_frame
        tf.clear()  # Clear default text
        p = tf.add_paragraph()
        p.text = subtitle
        p.font.size = Pt(18)
        p.alignment = PP_ALIGN.CENTER

        # Add body bullets to content placeholder
        for bullet_elem in body.findall('bullet'):
            p = tf.add_paragraph()
            p.text = bullet_elem.text
            p.level = 0  # Bullet level
            p.font.size = Pt(14)

        # Handle diagram: Save Mermaid code to file
        if diagram_code:
            diagram_file = f"diagram_{slide_id}.txt"
            with open(diagram_file, 'w', encoding='utf-8') as f:
                f.write(diagram_code)
            print(f"Saved Mermaid code for slide {slide_id} to {diagram_file}")

            # Check for diagram PNG and add to slide if it exists
            image_path = f"diagrams/diagram_{slide_id}.png"
            if os.path.exists(image_path):
                left = Inches(1)
                top = Inches(3)
                slide.shapes.add_picture(image_path, left, top, width=Inches(6))
            else:
                print(f"Warning: Image {image_path} not found. Please convert Mermaid code to PNG.")

    # Save the presentation
    prs.save('course.pptx')
    print("Presentation saved as 'course.pptx'")

except FileNotFoundError:
    print(f"Error: Could not find {xml_path}. Please ensure the file exists in the specified directory.")
except ET.ParseError as e:
    print(f"Error: XML parsing failed - {e}. Check slides_data.xml for syntax errors.")
except Exception as e:
    print(f"Unexpected error: {e}")